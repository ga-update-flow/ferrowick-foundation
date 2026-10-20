"""Generate every downloadable document for the Ferrowick Foundation demo site.

Formats mirror the reference site's document mix:
  * text PDFs printed from a grant portal (applications)
  * "scanned" image-only PDFs with no text layer (one sample application, CI guidance)
  * legacy Word .doc worksheets and Excel .xls/.xlsm workbooks (via LibreOffice)
  * report-style PDFs (biennial reports, highlights, research reports, job descriptions, slide deck)
  * an MP4 information-session video built from the slide deck

Usage: python3 scripts/docs/make_docs.py   (needs reportlab, pypdfium2, Pillow, python-docx, openpyxl,
LibreOffice `soffice` and `ffmpeg` on PATH). Reads build/documents.json written by `npm run docs`.
"""
import io
import json
import os
import random
import shutil
import subprocess
import sys
import tempfile

from PIL import Image, ImageFilter
import pypdfium2
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)
from reportlab.pdfgen import canvas as rl_canvas

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'src', 'documents')
REG = json.load(open(os.path.join(ROOT, 'build', 'documents.json')))
DOCS = REG['documents']
NAVY = colors.HexColor('#234E70')
BRASS = colors.HexColor('#C9963F')
ORG = 'The Ferrowick Foundation'


def out_path(href):
    p = os.path.join(OUT, href.lstrip('/'))
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


ss = getSampleStyleSheet()
S = {
    'title': ParagraphStyle('t', parent=ss['Title'], fontName='Helvetica', fontSize=24, leading=30, alignment=0, spaceAfter=6),
    'h1': ParagraphStyle('h1', parent=ss['Heading1'], fontName='Helvetica-Oblique', fontSize=19, leading=24, spaceBefore=10, spaceAfter=8),
    'h2': ParagraphStyle('h2', parent=ss['Heading2'], fontName='Helvetica-Bold', fontSize=13.5, leading=17, spaceBefore=10, spaceAfter=2),
    'h3': ParagraphStyle('h3', parent=ss['Heading3'], fontName='Helvetica-Bold', fontSize=12, leading=15, spaceBefore=8, spaceAfter=2, textColor=NAVY),
    'body': ParagraphStyle('b', parent=ss['BodyText'], fontName='Helvetica', fontSize=10.5, leading=14.5, spaceAfter=6),
    'small': ParagraphStyle('s', parent=ss['BodyText'], fontName='Helvetica-Oblique', fontSize=8.5, leading=11, textColor=colors.HexColor('#444444')),
    'guid': ParagraphStyle('g', parent=ss['BodyText'], fontName='Helvetica', fontSize=9.5, leading=12.5, textColor=colors.HexColor('#333333')),
    'choice': ParagraphStyle('c', parent=ss['BodyText'], fontName='Helvetica', fontSize=10, leading=13, leftIndent=18),
    'center': ParagraphStyle('ce', parent=ss['BodyText'], fontName='Helvetica', fontSize=10.5, leading=14, alignment=TA_CENTER),
    'cover': ParagraphStyle('cv', parent=ss['Title'], fontName='Times-Roman', fontSize=34, leading=40, textColor=colors.white, alignment=0),
}


# --------------------------------------------------------------------------- portal-style applications
def q(text, limit=None, guidance=None, choices=None, required=True):
    items = [Paragraph(text + ('*' if required else ''), S['h2'])]
    if guidance:
        items.append(Paragraph(guidance, S['guid']))
    if choices:
        items.append(Paragraph('Choices', S['small']))
        items += [Paragraph(c, S['choice']) for c in choices]
    if limit:
        items.append(Paragraph(f'Character Limit: {limit}', S['small']))
    return KeepTogether(items + [Spacer(1, 4)])


ADMIN = lambda support, grant=('One Time', 'Installment', 'Matching Grant', 'Director/Employee Matching Gift'): [
    Paragraph('Admin Only', S['h1']),
    q('Meeting Date', 10, required=False),
    q('Site Visit:', choices=['Yes', 'No', 'Not Necessary']),
    q('Type of Support:', choices=support),
    q('Grant Type', choices=list(grant)),
    q('Amount Recommended', 20),
    q('Source of Funding:', choices=['The Ferrowick Foundation Funds', 'Donor-Advised Partner Funds', 'Other Funds']),
    q('Notes:', 10000, required=False),
]

def org_info(long_limit=3500):
    """Organization questions shared by several applications; `long_limit` applies to A, B, C and F."""
    return [
    q('A. What is the purpose of your organization and its primary programs and services?', long_limit),
    q('B. What does your organization do well? Please highlight recent achievements.', long_limit),
    q('C. Apart from funding, what are the greatest challenges facing your organization or the people you serve?', long_limit),
    q('D. How does your organization plan for financial sustainability?', 2500),
    q('E. How does your organization address diversity, equity and inclusion?', 2500),
    q('F. Briefly describe partnerships or collaborations in which your organization takes part:', long_limit),
    q('G. Year established', 10),
    q('H. Number of full-time staff', 25),
    q('I. Number of part-time staff', 25),
    q('J. Total number of people served annually (unduplicated)', 25),
    q('K. Annual organizational budget', 20),
    ]


ORG_INFO = org_info()
FIN_INFO = [
    q('L. List total actual operating income and expenses for the last three completed fiscal years.', 1000,
      guidance="Use figures from your organization's IRS Form 990. Example: 2023, Income: $850,000, Expense: $830,000, Difference: $20,000"),
    q('M. Has your organization ended any of the past three fiscal years with a deficit?', choices=['Yes', 'No']),
    q('N. Does the organization currently have an operating reserve?', choices=['Yes', 'No']),
    q('O. Does the organization have an endowment or other funds not included in the annual budget?', choices=['Yes', 'No']),
    q('P. List all foundation and corporate supporters and all other sources of income, with amounts:', 4000,
      guidance='Include sources of income for the current and most recent fiscal year only.'),
    q("Q. Name of the organization's chief executive:", 100),
    q("R. Chief executive's title:", 100),
    q("S. Chief executive's email address:", 100),
    Paragraph('Deficit Information', S['h3']),
    q('Please explain the deficit and describe how the organization has addressed it.', 2000),
    Paragraph('Operating Reserve Information', S['h3']),
    q('What is the amount of the operating reserve?', 20),
    q('How many months of operating expenses does the reserve represent?', 20),
    Paragraph('Endowment Information', S['h3']),
    q('What is the current balance of those funds?', 20),
]
FOCUS_Q = q('A. Which Foundation Focus Area and target outcome does your project align with?',
            guidance='Guidance: Choose your Focus Area from the first menu, then select the ONE target outcome you expect to achieve. (If you are unsure, see The Ferrowick Foundation Guidelines.)',
            choices=[f['label'] for f in REG['focus']])
DOCS_SECTION = lambda extra=(): [
    Paragraph('Documentation', S['h1']),
    Paragraph('Please upload the following documents:', S['body']),
    q("Cover letter signed by the organization's CEO or Executive Director", guidance="Address the letter to the Foundation's Executive Director and include anything not covered in the application.") ,
    Paragraph('File Size Limit: 2 MB', S['small']),
    q("List of the Board of Directors or the organization's leaders"),
    Paragraph('File Size Limit: 3 MB', S['small']),
    q('Current organizational budget'), Paragraph('File Size Limit: 2 MB', S['small']),
    q('Project Budget Form', guidance='Please complete the Project Budget Template and upload it here.'),
    Paragraph('File Size Limit: 2 MB', S['small']),
    *extra,
    Paragraph('Authorized Signature', S['h1']),
    Paragraph('By typing your name, title and the date below, you confirm that you are authorized to submit this application on behalf of your organization.', S['body']),
    q('Name:', 75), q('Title:', 75), q('Date:', 10),
]

APPLICATIONS = {
    'orgStrength': ('Organizational Strength Application', '10 March 2022', [
        *ADMIN(['Project Support', 'General Operating Support', 'Strategic Planning Support', 'Technical Assistance Support', 'Technology Upgrade Support', 'Emergency Support', 'Director/Employee Matching Gift', 'Other Support']),
        Paragraph('I. Organization Information', S['h1']), *ORG_INFO, *FIN_INFO,
        Paragraph('II. Project Description', S['h1']),
        q('A. What issue will this project address for your organization and/or the people you serve?', 3500),
        q('B. How will your organization build its internal capacity to address this issue?', 3500),
        q('C. What activities are part of the project? Please include a timeline.', 4000),
        q('D. Project start date', guidance='Select a date with the calendar icon.'),
        q('E. Project end date', guidance='Select a date with the calendar icon.'),
        q('F. What resources does the organization already have that will be directed to this effort?', 2500),
        q('G. Apart from funding, what do you need to succeed (materials, staffing, expertise, etc.)?', 2500),
        q('H. Who will lead the work and who will assist (staff, consultants, project partners)?', 2500),
        q('I. If a consultant will be engaged, upload the proposed scope of work.', guidance='File Size Limit: 3 MB', required=False),
        Paragraph('III. Project Anticipated Results', S['h1']), FOCUS_Q,
        q('B. Choose the ONE area in which the new capacity will produce results:', choices=['Increased number of people served', 'Increased revenue', 'Decreased expenses', 'Decreased time to achieve results']),
        q('C. What short-term and long-term results do you anticipate for your organization?', 3500,
          guidance='A worksheet to help you write a clear, measurable and verifiable results statement is available on the Application Assistance page: Organizational Strength Results Worksheet.'),
        q('D. List the milestones you expect to reach during the grant period and when.', 3500,
          guidance='See the Organizational Strength Milestones Workbook on the Application Assistance page.'),
        q('E. How will you track and verify progress?', 2500),
        Paragraph('IV. Budget', S['h1']),
        q('A. Total project budget', 20), q('B. Amount requested', 20),
        q('C. Other sources of support for this project, committed or pending, with amounts', 2000),
        *DOCS_SECTION(),
    ]),
    'facilities': ('Facilities & Equipment Application', '10 March 2022', [
        *ADMIN(['Capital Support', 'Challenge Grant/Capital Support', 'Other Support']),
        Paragraph('I. Organization Information', S['h1']), *ORG_INFO, *FIN_INFO,
        Paragraph('II. Project Description', S['h1']),
        q('A. Describe the proposed facility or equipment project.', 3500),
        q('B. What issue will the project address?', 3500),
        q('C. What activities are part of the project? Please include a timeline.', 4000),
        q('D. Project start date', guidance='Select a date with the calendar icon.'),
        q('E. Project end date', guidance='Select a date with the calendar icon.'),
        q('F. What resources does the organization already have that will be directed to this effort?', 2500, required=False),
        q('G. Apart from funding, what does the project need to succeed?', 2500, required=False),
        q('H. Who will lead the work and who will assist (staff, contractors, project partners)?', 2500),
        q('I. Will the project require permits or approvals? If so, describe their status.', 2000),
        Paragraph('III. Project Anticipated Results', S['h1']), FOCUS_Q,
        q('B. What short-term and long-term results do you anticipate for your organization?', 3500,
          guidance='Note: This question concerns your organization. See the Facilities & Equipment Results Worksheet on the Application Assistance page.'),
        q('C. What short-term and long-term results do you anticipate for the people you serve?', 3500),
        q('D. List the milestones you expect to reach and when.', 3500),
        Paragraph('IV. Budget', S['h1']),
        q('A. Total project cost', 20), q('B. Amount requested', 20),
        q('C. Proposed payment schedule (by fiscal year)', 1000),
        q('D. Other sources of support for this project, committed or pending, with amounts', 2000),
        *DOCS_SECTION([q('Contractor estimates or vendor quotes'), Paragraph('File Size Limit: 5 MB', S['small'])]),
    ]),
    'rapid': ('Rapid Opportunity', '14 November 2022', [
        *ADMIN(['Project Support', 'General Operating Support', 'Capital Support', 'Emergency Support', 'Director/Employee Matching Gift', 'Other Support'], ('One Time', 'Installment', 'Matching Grant')),
        FOCUS_Q,
        Paragraph('I. Organization Information', S['h1']),
        q('Is your organization recognized by the IRS as a 501(c)(3) organization?', choices=['Yes', 'No']),
        q('Name of the CEO or head of the organization', 100), q('Title of the CEO or head of the organization', 100),
        q('Email address of the CEO or head of the organization', 100), q('Phone number of the CEO or head of the organization', 25),
        q('What is the purpose of your organization and its primary programs and services?', 3500),
        q('Year established', 10), q('Number of staff and/or volunteers', 25), q('Annual organizational budget', 20),
        q("Social media accounts other than the organization's website", 250, guidance='Example: Facebook, Instagram, blog', required=False),
        Paragraph('II. Project Description', S['h1']),
        q('How will the funds be used?', 250, guidance='(In 20 words or fewer.)'),
        q('Amount requested', 20),
        q('Project description and activities', 4000, guidance='Describe the problem you will address, the people you will work with, the services or activities you will provide and where, when the work will begin and end, and anything else you would like us to know.'),
        q('As a result of the project, what will be different for participants?', 3500),
        Paragraph('III. Documentation', S['h1']),
        *DOCS_SECTION()[2:],
    ]),
    'collective': ('Collective Impact Application', '02 June 2021', [
        *ADMIN(['Project Support', 'General Operating Support', 'Capital Support', 'Challenge Grant/Capital Support', 'Technical Assistance Support', 'Recovery Funds Support', 'Emergency Support', 'Other Support']),
        q('Emergency Funding Focus Area', choices=[f['label'] for f in REG['focus']], required=False),
        Paragraph('I. Organization Information', S['h1']),
        q("1. What is your organization's mission and vision?", 1500),
        q("2. Organization's current programs and services", 3500, guidance='Emphasize the programs and services related to the collective impact effort.'),
        q('3. What strengths and expertise does your organization bring?', 3500, guidance='What does your organization do well? This may include: Knowledge – insight into participant barriers and opportunities; Skills – what your organization knows how to do; Experience – length of time and track record in the proposed focus area.'),
        q('4. Plan for financial viability', 2500, guidance='The Ferrowick Foundation Board of Directors invests in financially viable organizations that can focus fully on carrying out their projects.'),
        q('5. Year established', 10), q('6. Number of full-time staff', 25), q('7. Number of part-time staff', 25),
        q('8. Annual operating budget', 20), q('9. Total number served annually (unduplicated)', 25),
        Paragraph('II. Request Snapshot', S['h1']),
        q('1. Project title', 100, guidance='Give the project a title of no more than ten words. The title should not include the name of the organization.'),
        q('2. Population served', guidance='Select the ONE population this grant will serve:', choices=['Youth', 'Older adults', 'People experiencing homelessness', 'People with physical disabilities', 'People with developmental disabilities', 'People with behavioral health needs', 'Disaster-affected residents', 'LGBTQ+ residents', 'People of color', 'Immigrants', 'General', 'Other']),
        q('3. Geographic area served', guidance='Select the geographic area this grant will serve:', choices=['Aquidneck Island (all)', 'Newport', 'Newport – North End', 'Middletown', 'Portsmouth', 'Newport County']),
        q('4. Diversity, Equity & Inclusion', 2500, guidance="For each category, give or estimate the percentage of people of color among Board members, staff and clients. Then describe how your organization's thinking about racial equity informs program design and delivery, and whether clients and other constituents take part in decision-making."),
        q('5. Overall project budget', 20, guidance='Enter the total budget for the project. This must equal the total on the required project budget attachment.'),
        q('6. Amount requested', 20),
        q('7. Project Budget Form', guidance='Please complete the Project Budget Template and upload it here.'),
        q('8. Focus Area and target outcome for the project', guidance='Choose your Focus Area, then select the ONE target outcome you expect to achieve. (If you are unsure, see The Ferrowick Foundation Guidelines.)', choices=[f['label'] for f in REG['focus']]),
        q('9. Project start date', guidance='Select a date with the calendar icon.'),
        q('10. Project end date', guidance='Select a date with the calendar icon.'),
        Paragraph('III. Collective Impact Alignment', S['h1']),
        q('1. Which collective impact initiative does this project advance?', choices=['Island Youth Prevention Coalition', 'Aquidneck Maternal and Infant Health Network', 'North End Early Learning Partnership', 'Island Child Wellness Collaborative', 'Pathways to Work Partnership', 'Equity Roundtable of Aquidneck Island']),
        q("2. Which of the initiative's shared measures will the project affect, and how?", 3500),
        q('3. Describe your role in the initiative (backbone, workgroup member, partner, etc.).', 2500),
        q('4. Describe the mutually reinforcing activities your project contributes.', 3500),
        q('5. How will results be reported to the initiative’s shared measurement system?', 2500),
        Paragraph('IV. Anticipated Results', S['h1']),
        q('1. Anticipated results statement', 3500, guidance='See the Collective Impact Proposal Guidance for examples of population-level and performance-level results.'),
        q('2. Milestones and timeline', 3500),
        *DOCS_SECTION([q('Letter of support from the initiative’s backbone organization'), Paragraph('File Size Limit: 2 MB', S['small'])]),
    ]),
}

PROGRAM_IMPACT = ('Program Impact Application', [
    Paragraph('I. Organization Information', S['h1']), *org_info(long_limit=4500),
    Paragraph('II. Project Description', S['h1']),
    q('A. Project title', 100, guidance='No more than ten words; do not include the organization name.'),
    q('B. What need or problem will the project address, and for whom?', 4500),
    q('C. Describe the project activities, where they will take place, and the timeline.', 4000),
    q('D. Project start date'), q('E. Project end date'),
    q('F. How many people will the project serve? Describe the core participants.', 1500),
    q('G. Which partners will be involved and what will each contribute?', 2500),
    Paragraph('III. Project Anticipated Results', S['h1']), FOCUS_Q,
    q('B. Anticipated results statement', 4500, guidance='Of the ___ participants we will serve, we expect ___ will achieve ___ by __/__/__. See the Program Impact Results Worksheet.'),
    q('C. List the milestones participants will reach and when.', 4500, guidance='See the Program Impact Milestones Workbook.'),
    q('D. What data will you collect to verify results?', 2500),
    Paragraph('IV. Budget', S['h1']),
    q('A. Total project budget', 20), q('B. Amount requested', 20),
    q('C. Other sources of support for this project, committed or pending, with amounts', 2000),
    *DOCS_SECTION(),
])


def portal_pdf(path, title, printed, story):
    def deco(c, d):
        c.saveState()
        c.setFont('Helvetica', 8.5)
        c.drawString(0.75 * inch, 10.55 * inch, 'Application')
        c.drawRightString(7.75 * inch, 10.55 * inch, ORG)
        c.drawString(0.75 * inch, 0.5 * inch, f'Printed On: {printed}')
        c.drawRightString(7.75 * inch, 0.5 * inch, f'{title}   {d.page}')
        c.restoreState()
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.85 * inch, rightMargin=0.85 * inch, topMargin=0.8 * inch,
                            bottomMargin=0.8 * inch, title=title, author=ORG, subject='Sample application (for review only)')
    doc.build([Paragraph(title, S['title']), Paragraph(ORG, S['body']), Spacer(1, 8)] + story, onFirstPage=deco, onLaterPages=deco)


def scanned_pdf(text_pdf_bytes, path, seed, title):
    """Rasterize a PDF and rebuild it from page images only, like a photocopied scan (no text layer)."""
    rnd = random.Random(seed)
    src = pypdfium2.PdfDocument(text_pdf_bytes)
    c = rl_canvas.Canvas(path, pagesize=letter)
    c.setTitle(title)
    c.setAuthor(ORG)
    for i in range(len(src)):
        im = src[i].render(scale=1.6).to_pil().convert('L')
        im = im.rotate(rnd.uniform(-0.7, 0.7), resample=Image.BICUBIC, fillcolor=255, expand=False)
        noise = Image.effect_noise(im.size, 18).point(lambda v: 255 if v > 110 else 235)
        im = Image.composite(im, noise, im.point(lambda v: 255 if v < 250 else 0)).filter(ImageFilter.GaussianBlur(0.35))
        buf = io.BytesIO()
        im.save(buf, 'JPEG', quality=62)
        buf.seek(0)
        from reportlab.lib.utils import ImageReader
        c.drawImage(ImageReader(buf), 0, 0, width=letter[0], height=letter[1])
        c.showPage()
    c.save()


def build_applications():
    for key, d in DOCS.items():
        if d['kind'] == 'application':
            title, printed, story = APPLICATIONS[d['type']]
            portal_pdf(out_path(d['href']), title, printed, story)
            print('pdf', d['href'])
    d = DOCS['appProgramImpact']
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.9 * inch, bottomMargin=0.9 * inch)

    def deco(c, dd):
        c.setFont('Helvetica', 8.5)
        c.drawString(0.9 * inch, 10.5 * inch, 'Application')
        c.drawRightString(7.6 * inch, 10.5 * inch, ORG)
        c.drawCentredString(4.25 * inch, 0.55 * inch, 'Program Impact Application')
    doc.build([Paragraph('Program Impact Application', S['title']), Spacer(1, 6)] + PROGRAM_IMPACT[1], onFirstPage=deco, onLaterPages=deco)
    scanned_pdf(buf.getvalue(), out_path(d['href']), 7, 'Program Impact Application')
    print('pdf (scanned)', d['href'])


# --------------------------------------------------------------------------- CI guidance (scanned)
def build_ci_guidance():
    d = DOCS['ciGuidance']
    buf = io.BytesIO()
    story = [Paragraph('The Ferrowick Foundation', S['h3']), Paragraph('Collective Impact Proposal Guidance', S['title']),
             Paragraph('February 14, 2019', S['body']), Spacer(1, 10),
             Paragraph('Purpose', S['h2']),
             Paragraph('This guidance is for organizations proposing programs or projects that advance a collective impact (CI) initiative on Aquidneck Island. CI proposals differ from other Foundation requests because they are judged both on the results the applicant will achieve and on how those results contribute to the population-level goals of a shared initiative.', S['body']),
             Paragraph('Before You Apply', S['h2']),
             Paragraph('Applicants must discuss a potential CI proposal with Foundation staff and with the backbone organization of the relevant initiative before beginning an application. Proposals are submitted through the Foundation’s online portal by invitation, using the Collective Impact Application.', S['body']),
             Paragraph('What a Strong Proposal Includes', S['h2'])]
    for t in ['A clear connection to the initiative’s common agenda and at least one shared measure.',
              'A description of the applicant’s role and the mutually reinforcing activities it will carry out.',
              'An anticipated results statement at the performance level (what will change for participants) and a description of the expected contribution to population-level change.',
              'A plan for reporting data to the initiative’s shared measurement system.',
              'A letter of support from the initiative’s backbone organization.']:
        story.append(Paragraph('• ' + t, S['body']))
    story += [Paragraph('Population-Level vs. Performance-Level Results', S['h2']),
              Paragraph('Population-level results describe conditions for everyone in a community (for example, the share of island third graders reading on grade level). Performance-level results describe what changes for the people a specific program serves (for example, 40 of 55 tutored students gain one reading level). The Foundation asks CI applicants to commit to performance-level results and to explain how they contribute to the population-level result.', S['body']),
              Paragraph('Review', S['h2']),
              Paragraph('CI proposals follow the Foundation’s regular Board calendar and deadlines. Staff may consult the initiative’s steering committee during review.', S['body']),
              PageBreak()]
    attachments = [
        ('Attachment A: Collective Impact Conditions Checklist', [['Condition', 'Evidence in proposal', 'Yes / No'], ['Common agenda', '', ''], ['Shared measurement', '', ''], ['Mutually reinforcing activities', '', ''], ['Continuous communication', '', ''], ['Backbone support', '', '']]),
        ('Attachment B: Proposal Outline', [['Section', 'Suggested length'], ['1. Initiative and common agenda', '1/2 page'], ['2. Your role and activities', '1 page'], ['3. Anticipated results', '1/2 page'], ['4. Shared measurement and reporting', '1/2 page'], ['5. Budget narrative', '1/2 page']]),
        ('Attachment C: Results Framework Template', [['Population result', 'Indicator', 'Program result', 'Performance measure', 'Data source'], ['', '', '', '', ''], ['', '', '', '', ''], ['', '', '', '', '']]),
        ('Attachment D: Project Budget Template', [['Line item', 'Total cost', 'Request to Foundation', 'Other sources'], ['Personnel', '', '', ''], ['Fringe', '', '', ''], ['Consultants', '', '', ''], ['Supplies', '', '', ''], ['Other', '', '', ''], ['TOTAL', '', '', '']]),
        ('Attachment E: Initiative Contacts', [['Initiative', 'Backbone organization'], ['Island Youth Prevention Coalition', 'Aquidneck Partnership for Wellness'], ['North End Early Learning Partnership', 'Island Learning Partnership'], ['Aquidneck Maternal and Infant Health Network', 'Island Community Health Workers'], ['Pathways to Work Partnership', 'Lantern Wharf Workforce Collaborative']]),
    ]
    for i, (title, rows) in enumerate(attachments):
        t = Table(rows, repeatRows=1, rowHeights=[22] + [44] * (len(rows) - 1))
        t.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.6, colors.black), ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'), ('FONTSIZE', (0, 0), (-1, -1), 9.5), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
        story += [Paragraph(title, S['h2']), Spacer(1, 8), t]
        if i < len(attachments) - 1:
            story.append(PageBreak())
    SimpleDocTemplate(buf, pagesize=letter, leftMargin=inch, rightMargin=inch).build(story)
    scanned_pdf(buf.getvalue(), out_path(d['href']), 19, 'Collective Impact Proposal Guidance')
    print('pdf (scanned)', d['href'])


# --------------------------------------------------------------------------- worksheets (.doc)
WS = {
    'programImpact': ('Program Impact Results Worksheet', 'participants', [
        'improving grades and school attendance', 'reducing conflict at home', 'building positive friendships with peers',
        'reducing alcohol or vaping use', 'getting and keeping an unsubsidized job', 'families who were homeless securing stable housing'],
        ['cutting failing grades from an average of two to none', 'arguments at home dropping from daily to once a week',
         'forming at least one friendship with a peer who values school', 'no vaping for at least 90 days',
         'earning at least $18 per hour and keeping the job for six months', 'keeping housing for at least twelve months']),
    'orgStrength': ('Organizational Strength Results Worksheet', 'organization', [
        'serving more people with the same staff', 'raising more unrestricted revenue', 'lowering the cost of a core service',
        'shortening the time it takes to enroll a new client', 'reducing staff turnover'],
        ['serving 120 families a month instead of 90', 'increasing individual giving from $40,000 to $65,000',
         'reducing the cost per meal from $6.10 to $5.25', 'cutting intake time from 90 minutes to 40 minutes', 'reducing annual turnover from 30% to 15%']),
    'facilities': ('Facilities & Equipment Results Worksheet', 'organization and participants', [
        'opening space that allows more people to be served', 'lowering operating costs through efficiency',
        'improving safety or accessibility', 'adding a new service that the old space could not support'],
        ['enrolling 24 additional children in the new classroom', 'reducing heating costs by 30%',
         'zero accessibility complaints after the ramp is installed', 'offering evening clinic hours three nights a week']),
}


def build_worksheets(tmp):
    from docx import Document
    from docx.shared import Pt, RGBColor
    for key, d in DOCS.items():
        if d['kind'] != 'worksheet':
            continue
        title, who, changes, amounts = WS[d['type']]
        doc = Document()
        st = doc.styles['Normal']
        st.font.name = 'Arial'
        st.font.size = Pt(10.5)
        h = doc.add_heading(title, 0)
        doc.add_paragraph(f'{ORG} · Use this worksheet to draft a clear, measurable and verifiable results statement for your application.')

        def step(label, text):
            p = doc.add_paragraph()
            r = p.add_run(label)
            r.bold = True
            r.font.color.rgb = RGBColor(0x23, 0x4E, 0x70)
            p.add_run(' – ' + text)

        def box(prompt):
            doc.add_paragraph(prompt).runs[0].italic = True
            t = doc.add_table(rows=1, cols=1)
            t.style = 'Table Grid'
            t.cell(0, 0).text = 'Enter your response here.\n\n\n'
        step('STEP ONE', f'Identify the kind of change you want to see for your {who}. Examples:')
        for c in changes:
            doc.add_paragraph(c, style='List Bullet')
        box(f'In the space below, list the kinds of change you want to see for your {who}.')
        step('STEP TWO', 'Specify how much change is needed for you to consider it a success. Examples:')
        for c in amounts:
            doc.add_paragraph(c, style='List Bullet')
        box('In the space below, describe how much change is needed for you to call it a success.')
        step('STEP THREE', 'Estimate how many will change, by how much, and by when.')
        doc.add_paragraph('Of the _____ participants we will serve, we expect _____ will achieve results by __/__/__')
        box('Enter your projections in the box below.')
        step('STEP FOUR', 'Combine the three steps above into a structured results statement, using a single-focus, menu or individualized approach.')
        doc.add_paragraph('Single focus: By June 30, 2026, 14 of the 20 adults we serve this year will get and keep a job paying at least $18 per hour for at least three months.')
        doc.add_paragraph('Menu: Of the 120 students enrolled in grades 6–7 at the participating middle school, 80 will improve in at least two of the following ways by June 15, 2026: fewer failing grades; fewer absences; one new positive peer friendship; no vaping for 90 days.')
        doc.add_paragraph('Individualized: 30 of the 40 families who secure housing will keep it for at least twelve months by meeting the financial and employment goals in their individual case plans by September 30, 2026.')
        box('Use the space below to write your results statement.')
        step('BACK TO YOUR APPLICATION', 'Copy your final results statement(s) into the appropriate fields of your online application. You may find that you expect results in more than one way, or in ways you did not first anticipate.')
        doc.add_paragraph('You may also save this file and print it for your records.')
        docx = os.path.join(tmp, os.path.basename(d['href']).replace('.doc', '.docx'))
        doc.save(docx)
        convert(docx, 'doc:MS Word 97', out_path(d['href']), tmp)


# --------------------------------------------------------------------------- workbooks (.xls / .xlsm)
WB_EXAMPLES = {
    'programImpact': ('Program Example', 'Program Impact Example\nYouth Employment: Summer Harbor Jobs', [
        ('Recruit through schools, libraries and youth programs; hold two information sessions.', 'At least 40 youth ages 16–19 apply.', 'Month 1–2'),
        ('Run a three-week job-readiness workshop series.', 'At least 32 youth complete the workshops.', 'Month 3'),
        ('Match youth with waterfront employers; provide transportation passes.', 'At least 28 youth start summer jobs.', 'Month 4'),
        ('Weekly check-ins with youth and supervisors.', 'At least 24 youth complete eight weeks of work.', 'Month 6'),
        ('Fall follow-up and referrals to school-year jobs.', 'At least 15 youth hold a school-year job.', 'Month 9–12')]),
    'orgStrength': ('Building Capacity Example', 'Organizational Strength Example\nAdopting a Client Data System', [
        ('Agree on common intake fields with program and intake staff; configure software.', 'System set up so all key staff can enter and view data.', 'Month 2'),
        ('Train staff; revise intake procedures.', 'Key staff show they can enter and use data in real time.', 'Month 4'),
        ('Use the new system for 30 days.', 'Intake completed in 60 minutes or less.', 'Month 5'),
        ('Use the new system for 90 more days.', 'Intake completed in 45 minutes or less.', 'Month 8'),
        ('Use the new system for 90 more days.', 'Intake completed in 30 minutes; enrollment in 15 minutes or less.', 'Month 11')]),
    'facilities': ('Capital or Technology Example', 'Facilities & Equipment Example\nExpanding a Teaching Kitchen', [
        ('Submit plans to the building department.', 'All required permits obtained.', 'Month 2'),
        ('Complete electrical, plumbing, ventilation and equipment installation.', 'Kitchen passes inspection.', 'Month 6'),
        ('Hire and train two instructors.', 'First class cohort enrolled.', 'Month 7'),
        ('Run three 8-week culinary training cohorts.', 'At least 30 adults complete training.', 'Month 8–14'),
        ('Job placement with island restaurants.', 'At least 20 graduates employed.', 'Month 15–18')]),
}


def build_workbooks(tmp):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wrap = Alignment(wrap_text=True, vertical='top')
    head = Font(bold=True, color='FFFFFF')
    fill = PatternFill('solid', fgColor='234E70')
    for key, d in DOCS.items():
        if d['kind'] != 'workbook':
            continue
        wb = Workbook()
        ws = wb.active
        ws.title = 'Instructions - READ FIRST'
        ws.column_dimensions['A'].width = 110
        rows = ['GENERAL INSTRUCTIONS',
                f'{d["title"]} · {ORG}',
                'Use this workbook to define the key progress points (milestones) you expect your organization or participants to reach, the steps you will take to get there, and when you expect each milestone to be achieved. Please review the example on the last tab before completing your own.',
                'Step 1 – Milestones & Timeline: Whatever the length of your project, list 3–5 progress points that show you are on the way to your result, so you can check interim progress and adjust. Click the STEP 1 tab below.',
                'Step 2 – Project Steps: For each milestone, summarize the critical activities that must happen for it to be reached. Use “if–then” logic: if we complete these activities, then participants or the organization will reach these milestones. Click the STEP 2 tab below.',
                'Step 3 – Copy into the Application: The STEP 3 tab combines your entries into a single statement you can copy into the application form.']
        for i, r in enumerate(rows, 1):
            ws.cell(i, 1, r).alignment = wrap
        ws['A1'].font = Font(bold=True, size=14)
        s1 = wb.create_sheet('Step 1')
        s1.column_dimensions['A'].width = 70
        s1.column_dimensions['B'].width = 30
        s1['A1'] = 'INSTRUCTIONS: Think about the result you want. What 3–5 progress points during the grant period will show you are on your way? List changes that can be easily confirmed.'
        s1['B1'] = 'List the date or timeframe for each milestone.'
        s1['A3'], s1['B3'] = 'Milestones', 'By When'
        for c in ('A3', 'B3'):
            s1[c].font, s1[c].fill = head, fill
        for i in range(5):
            s1.cell(4 + i, 1, f'Enter Milestone {i + 1} here.' if i == 0 else f'Enter Milestone {i + 1} here or delete these contents.')
            s1.cell(4 + i, 2, 'Enter date/timeframe here.')
        s2 = wb.create_sheet('Step 2')
        for col, w in zip('ABCD', (8, 60, 50, 20)):
            s2.column_dimensions[col].width = w
        s2['B1'] = 'INSTRUCTIONS: For each milestone, list the activities needed to reach it.'
        for c, v in zip(('A3', 'B3', 'C3', 'D3'), ('Step', 'Critical Steps', 'Milestones', 'By When')):
            s2[c] = v
            s2[c].font, s2[c].fill = head, fill
        for i in range(5):
            r = 4 + i
            s2.cell(r, 1, i + 1)
            s2.cell(r, 2, 'Enter steps to achieve this milestone.' if i == 0 else ' ')
            s2.cell(r, 3, f"='Step 1'!A{r}")
            s2.cell(r, 4, f"='Step 1'!B{r}")
        s3 = wb.create_sheet('Step 3')
        s3.column_dimensions['A'].width = 120
        s3['A1'] = 'Your final statement is below. Copy and paste it into your application.'
        s3['A3'] = "=\"Step 1: \"&'Step 2'!B4&\" Milestone: \"&'Step 2'!C4&\" (\"&'Step 2'!D4&\")\"&CHAR(10)&\"Step 2: \"&'Step 2'!B5&\" Milestone: \"&'Step 2'!C5&\" (\"&'Step 2'!D5&\")\""
        s3['A3'].alignment = wrap
        name, label, ex = WB_EXAMPLES[d['type']]
        se = wb.create_sheet(name)
        for col, w in zip('ABCD', (8, 60, 50, 16)):
            se.column_dimensions[col].width = w
        se['B1'] = label
        se['B1'].alignment = wrap
        for c, v in zip(('A2', 'B2', 'C2', 'D2'), ('Step', 'Critical Steps', 'Milestones', 'By When')):
            se[c] = v
            se[c].font, se[c].fill = head, fill
        for i, (a, b, c) in enumerate(ex):
            for j, v in enumerate((i + 1, a, b, c)):
                se.cell(3 + i, 1 + j, v).alignment = wrap
        xlsx = os.path.join(tmp, os.path.basename(d['href']).rsplit('.', 1)[0] + '.xlsx')
        wb.save(xlsx)
        if d['href'].endswith('.xlsm'):
            convert(xlsx, 'xlsm:Calc MS Excel 2007 VBA XML', out_path(d['href']), tmp)
        else:
            convert(xlsx, 'xls:MS Excel 97', out_path(d['href']), tmp)


def convert(src, fmt, dest, tmp):
    ext = fmt.split(':')[0]
    subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', tmp, src], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    produced = os.path.join(tmp, os.path.basename(src).rsplit('.', 1)[0] + '.' + ext)
    shutil.move(produced, dest)
    print(ext, os.path.relpath(dest, OUT))


# --------------------------------------------------------------------------- report PDFs
def report_doc(path, title, story, pagesize=letter):
    def deco(c, d):
        c.saveState()
        c.setFillColor(BRASS)
        c.rect(0, pagesize[1] - 10, pagesize[0], 10, fill=1, stroke=0)
        c.setFont('Helvetica', 8)
        c.setFillColor(colors.HexColor('#555555'))
        c.drawString(0.75 * inch, 0.5 * inch, f'{ORG} · {title}')
        c.drawRightString(pagesize[0] - 0.75 * inch, 0.5 * inch, str(d.page))
        c.restoreState()
    SimpleDocTemplate(path, pagesize=pagesize, title=title, author=ORG, leftMargin=0.9 * inch, rightMargin=0.9 * inch,
                      topMargin=0.9 * inch, bottomMargin=0.85 * inch).build(story, onFirstPage=deco, onLaterPages=deco)


def cover(title, subtitle):
    t = Table([[Paragraph(title, S['cover'])], [Paragraph(f'<font color="#F3D28C" size="15">{subtitle}</font>', S['body'])]],
              colWidths=[6.7 * inch], rowHeights=[3.4 * inch, 1.2 * inch])
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), NAVY), ('VALIGN', (0, 0), (-1, 0), 'BOTTOM'), ('LEFTPADDING', (0, 0), (-1, -1), 24)]))
    return [Spacer(1, 0.6 * inch), t, Spacer(1, 0.3 * inch), Paragraph(ORG, S['h3']),
            Paragraph('75 Lantern Wharf Way, Suite 2C · Middletown, RI 02842 · info@example.org', S['body']), PageBreak()]


def money(n):
    return '${:,.0f}'.format(n)


def build_biennials():
    chairs = {2000: 'Cornelius A. Brightman', 2010: 'Harold Venner', 2018: 'Margaret Halloran Price'}
    for i, r in enumerate(REG['biennial']):
        y1, y2 = [int(x) for x in r['period'].split('/')]
        rnd = random.Random(y1)
        chair = [v for k, v in sorted(chairs.items()) if k <= y1][-1]
        ed = 'Rosalind Atherly, Executive Director' if y1 < 2020 else 'Rosemary Quintal, Acting Executive Director'
        scale = 1 + (y1 - 2000) * 0.06
        areas = ['Education', 'Health', 'Community Services', 'Arts & Culture']
        amounts = [int(rnd.uniform(0.6, 1.1) * 900000 * scale / 1000) * 1000 for _ in areas]
        counts = [rnd.randint(18, 46) for _ in areas]
        story = cover(f'{r["period"]} Biennial Report', f'Fiscal Years {y1} and {y2}')
        story += [Paragraph('Letter from the Chair and Executive Director', S['h1']),
                  Paragraph(f'Dear Friends, In {y1} and {y2} the Ferrowick Foundation continued its commitment to improving the quality of life on Aquidneck Island, especially for residents with the fewest resources. Over these two years the Board approved {sum(counts)} grants totaling {money(sum(amounts))} across our four focus areas.', S['body']),
                  Paragraph('We are grateful to the nonprofit leaders, volunteers and partners whose work these grants support, and to the residents who shared their ideas and concerns with us. This report summarizes our grantmaking and finances for the period.', S['body']),
                  Paragraph(f'{chair}, Chair<br/>{ed}', S['body']),
                  Paragraph('Grants Approved by Focus Area', S['h1'])]
        t = Table([['Focus Area', 'Number of grants', 'Amount approved']] + [[a, str(c), money(m)] for a, c, m in zip(areas, counts, amounts)] + [['Total', str(sum(counts)), money(sum(amounts))]],
                  colWidths=[2.8 * inch, 1.6 * inch, 1.8 * inch])
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), NAVY), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'), ('GRID', (0, 0), (-1, -1), 0.4, colors.grey), ('ALIGN', (1, 0), (-1, -1), 'RIGHT')]))
        story += [t, Paragraph('Financial Summary', S['h1'])]
        assets = int(62_000_000 * (1 + (y1 - 2000) * 0.025) * rnd.uniform(0.92, 1.05) / 1000) * 1000
        f = Table([['', f'FY {y1}', f'FY {y2}'], ['Total assets (year end)', money(assets), money(assets * rnd.uniform(0.97, 1.06))],
                   ['Grants paid', money(sum(amounts) * 0.47), money(sum(amounts) * 0.53)], ['Program and administrative expenses', money(410000 * scale), money(425000 * scale)]],
                  colWidths=[3 * inch, 1.6 * inch, 1.6 * inch])
        f.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 0.4, colors.grey), ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EEF1F4')), ('ALIGN', (1, 0), (-1, -1), 'RIGHT')]))
        story += [f, Paragraph('Figures are unaudited summaries prepared for this report.', S['small']), PageBreak(),
                  Paragraph('Selected Grants', S['h1'])]
        for a in areas:
            story.append(Paragraph(a, S['h3']))
            for k in range(3):
                story.append(Paragraph(f'• {rnd.choice(SELECTED[a])}', S['body']))
        doc_path = out_path(r['href'])
        report_doc(doc_path, f'{r["period"]} Biennial Report', story)
        print('pdf', r['href'])


SELECTED = {
    'Education': ['Reading coaches for kindergarten classrooms in two Newport elementary schools', 'An after-school homework center for middle school students', 'English classes for adult learners at the public library', 'A college-access advising program for first-generation students', 'Summer learning camps for rising first graders'],
    'Health': ['A community health worker serving families in the North End', 'Free fitness classes for older adults', 'Youth substance-use prevention programming in island middle schools', 'Sliding-scale counseling for uninsured residents', 'School garden and nutrition lessons'],
    'Community Services': ['Emergency heating assistance for households with low incomes', 'Expanded weekend food pantry hours', 'Eviction-prevention outreach for older adults', 'A youth civic leadership program', 'Board development consulting for small nonprofits'],
    'Arts & Culture': ['String instruction residencies in elementary schools', 'Free family days at a community arts center', 'Maritime-history field studies for middle school students', 'A shared marketing platform for small arts groups', 'A teen summer theater intensive'],
}


def build_simple_reports():
    d = DOCS['highlights']
    story = cover('Highlights 1999–2025', 'Twenty-five years of grantmaking on Aquidneck Island')
    timeline = [('1999', 'Ferrowick Savings & Trust converts to a stock bank and endows The Ferrowick Foundation. Cornelius A. Brightman becomes founding Chairman.'),
                ('2000', 'Rosalind “Roz” Atherly is named the Foundation’s first executive director. First grants are awarded.'),
                ('2004', 'Cumulative grants pass $10 million.'), ('2008', 'Results-based grantmaking adopted for all requests.'),
                ('2010', 'Harold Venner becomes Chairman. Vision for Aquidneck Island community engagement project is published.'),
                ('2012', 'The Foundation convenes the Island Youth Prevention Coalition, its first collective impact initiative.'),
                ('2018', 'North End well-being study commissioned. Margaret Halloran Price becomes Chair.'),
                ('2020', 'Emergency grants respond to the COVID-19 pandemic.'),
                ('2021', 'Roz Atherly passes away; the Board establishes the Rosalind Atherly Fellowship in Social Work.'),
                ('2022', 'Rosemary Quintal is appointed Executive Director. First Fellowship cohort begins.'),
                ('2025', 'Keystone Matching Grant Program launched. Cumulative grants exceed $58 million.')]
    t = Table([[Paragraph(f'<b>{y}</b>', S['body']), Paragraph(x, S['body'])] for y, x in timeline], colWidths=[0.9 * inch, 5.6 * inch])
    t.setStyle(TableStyle([('LINEBELOW', (0, 0), (-1, -1), 0.3, colors.lightgrey), ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
    report_doc(out_path(d['href']), 'Highlights 1999–2025', story + [Paragraph('Milestones', S['h1']), t])
    print('pdf', d['href'])

    d = DOCS['visionReport']
    story = cover('Vision for Aquidneck Island', 'Islanders Building Their Future · 2010')
    story += [Paragraph('Summary', S['h1']),
              Paragraph('In 2009 and 2010, more than 700 people who live and work on Aquidneck Island answered a simple question: what can islanders do to improve their own lives and the lives of their neighbors? Community conversations were held in libraries, schools, houses of worship and workplaces in Newport, Middletown and Portsmouth.', S['body']),
              Paragraph('Themes', S['h1'])]
    for th in ['Ready to learn: every child enters school prepared and supported', 'Healthy neighborhoods: access to care, healthy food and safe places to be active', 'Good jobs year-round: reducing the effects of seasonal employment', 'Connected community: transportation, civic participation and belonging', 'A living culture: arts and history that include every resident']:
        story.append(Paragraph('• ' + th, S['body']))
    story += [Paragraph('This report was prepared with support from The Ferrowick Foundation.', S['small'])]
    report_doc(out_path(d['href']), 'Vision for Aquidneck Island', story)
    print('pdf', d['href'])

    d = DOCS['wellbeingReport']
    story = cover('The North End of Newport', 'Community-Driven Solutions for Child and Family Well-Being · 2018')
    story += [Paragraph('About This Study', S['h1']),
              Paragraph('Commissioned by The Ferrowick Foundation and prepared by the Narragansett Bay Policy Lab, this study compares conditions for children and families in Newport’s North End with the rest of Aquidneck Island and the state across eight areas of well-being.', S['body']),
              Paragraph('Indicators (illustrative)', S['h1'])]
    rows = [['Area', 'North End', 'Rest of island'], ['Children in families below poverty', '31%', '9%'], ['Third graders meeting reading expectations', '29%', '57%'],
            ['Households spending >50% of income on rent', '38%', '17%'], ['Adults without a usual source of care', '19%', '8%'], ['Chronic absenteeism (K–12)', '27%', '13%']]
    t = Table(rows, colWidths=[3.6 * inch, 1.3 * inch, 1.3 * inch])
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), NAVY), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white), ('GRID', (0, 0), (-1, -1), 0.4, colors.grey)]))
    story += [t, Paragraph('Sources: state and municipal administrative data, 2016–2017; resident survey, 2018.', S['small'])]
    report_doc(out_path(d['href']), 'North End Well-Being Study', story)
    print('pdf', d['href'])


def build_jobs():
    jobs = {
        'intern': ('2024 Student Intern', 'Part-time, academic year (approximately 12 hours per week)', [
            ('About the Foundation', 'The Ferrowick Foundation works to improve the quality of life on Aquidneck Island, especially for residents with the fewest resources.'),
            ('Responsibilities', '• Support grant review preparation and data entry in the grants management system<br/>• Help organize convenings and fellowship learning sessions<br/>• Research community data for program staff<br/>• Assist with website and newsletter updates'),
            ('Qualifications', '• Current undergraduate or graduate student<br/>• Strong writing and organizational skills<br/>• Interest in nonprofits and community development'),
            ('To Apply', 'Send a cover letter and résumé to info@example.org with “Intern” in the subject line. Applications are reviewed on a rolling basis.')]),
        'po': ('Program Officer', 'Full-time, exempt', [
            ('About the Position', 'The Program Officer works with applicants and grantees across all four Focus Areas, guiding organizations through the Foundation’s results-based application process and monitoring grant progress.'),
            ('Responsibilities', '• Respond to inquiries from potential applicants and advise on alignment with Foundation targets<br/>• Review proposals, conduct site visits and prepare recommendations for the Board<br/>• Review progress and final reports and track results<br/>• Represent the Foundation in community partnerships and collective impact initiatives'),
            ('Qualifications', '• Five or more years of experience in nonprofit programs, philanthropy or public service<br/>• Knowledge of outcome measurement<br/>• Excellent communication skills; familiarity with Aquidneck Island communities preferred'),
            ('Compensation', 'Competitive salary and benefits commensurate with experience.'),
            ('To Apply', 'Send a cover letter and résumé to info@example.org with “Program Officer” in the subject line.')]),
    }
    for key, d in DOCS.items():
        if d['kind'] != 'job':
            continue
        title, kind, sections = jobs[d['type']]
        story = [Paragraph(ORG, S['h3']), Paragraph(title, S['title']), Paragraph(kind, S['body']), Spacer(1, 8)]
        for h, b in sections:
            story += [Paragraph(h, S['h2']), Paragraph(b, S['body'])]
        report_doc(out_path(d['href']), title, story)
        print('pdf', d['href'])


SLIDES = [
    ('Rosalind Atherly Fellowship in Social Work', 'Information Session · 2026–2027 Program Year\nThe Ferrowick Foundation'),
    ('Why this Fellowship?', 'Most social work field placements are unpaid.\nThe Fellowship pays island nonprofits so they can pay Fellows for internship hours.'),
    ('Who can apply?', '• Aquidneck Island residents\n• Accepted into a BSW or MSW program (full or part time)\n• No other source paying for internship hours\n• Committed to a social work career on the island'),
    ('What Fellows do', '• Full academic year placement at an island nonprofit (Fall 2026 – Spring 2027)\n• Monthly learning sessions with local social workers\n• Annual conference and networking events\n• Final report and experience survey'),
    ('How to apply', '• Review the How To Apply page and this presentation\n• Complete the online application\n• Deadline: February 27, 2026\n• We anticipate 2–3 Fellows for 2026–2027'),
    ('Questions?', 'info@example.org\nTobias Mendel-Ayers, Special Projects Manager'),
]


def build_fellowship_media(tmp):
    d = DOCS['fellowshipDeck']
    path = out_path(d['href'])
    W, H = landscape(letter)
    c = rl_canvas.Canvas(path, pagesize=(W, H))
    c.setTitle('Rosalind Atherly Fellowship – Information Session')
    c.setAuthor(ORG)
    for i, (title, body) in enumerate(SLIDES):
        c.setFillColor(NAVY if i in (0, len(SLIDES) - 1) else colors.white)
        c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(BRASS)
        c.rect(0, H - 18, W, 18, fill=1, stroke=0)
        fg = colors.white if i in (0, len(SLIDES) - 1) else NAVY
        c.setFillColor(fg)
        c.setFont('Times-Bold', 34)
        c.drawString(60, H - 110, title)
        c.setFont('Helvetica', 19)
        for j, line in enumerate(body.split('\n')):
            c.drawString(64, H - 180 - j * 34, line)
        c.setFont('Helvetica', 9)
        c.drawRightString(W - 40, 24, f'{ORG} · {i + 1}/{len(SLIDES)}')
        c.showPage()
    c.save()
    print('pdf', d['href'])
    # video: each slide shown for 5 seconds
    pdf = pypdfium2.PdfDocument(path)
    frames = []
    for i in range(len(pdf)):
        f = os.path.join(tmp, f'slide{i:02d}.png')
        im = pdf[i].render(scale=1.2).to_pil().convert('RGB')
        im = im.resize((1280, 990)).crop((0, 0, 1280, 990))
        canvas_im = Image.new('RGB', (1280, 990), 'white')
        canvas_im.paste(im)
        canvas_im.save(f)
        frames.append(f)
    lst = os.path.join(tmp, 'frames.txt')
    with open(lst, 'w') as fh:
        for f in frames:
            fh.write(f"file '{f}'\nduration 5\n")
        fh.write(f"file '{frames[-1]}'\n")
    v = DOCS['fellowshipVideo']
    subprocess.run(['ffmpeg', '-y', '-loglevel', 'error', '-f', 'concat', '-safe', '0', '-i', lst, '-vf', 'scale=1280:990,format=yuv420p', '-r', '25',
                    '-c:v', 'libx264', '-movflags', '+faststart', out_path(v['href'])], check=True)
    print('mp4', v['href'])


def main():
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT)
    with tempfile.TemporaryDirectory() as tmp:
        build_applications()
        build_ci_guidance()
        build_worksheets(tmp)
        build_workbooks(tmp)
        build_biennials()
        build_simple_reports()
        build_jobs()
        build_fellowship_media(tmp)


if __name__ == '__main__':
    sys.exit(main())
